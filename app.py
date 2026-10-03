import os, io, json, re
from datetime import datetime
import requests
from flask import Flask, request, jsonify, send_file
from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import create_engine, MetaData, Table, Column, Integer, String, Text, insert, select

app = Flask(__name__, static_folder="static")
DB_URL = os.environ.get("DATABASE_URL", "sqlite:///data.db").replace("postgres://", "postgresql://", 1)
engine = create_engine(DB_URL, pool_pre_ping=True)
md = MetaData()
plans = Table("plans", md, Column("id", Integer, primary_key=True), Column("created", String(32)),
              Column("name", String(200)), Column("payload", Text))
md.create_all(engine)

TYPE_W = {"emergency": 100, "education": 70, "home": 60, "retirement": 60, "other": 40}
WE = {"neg_cf": "Negative monthly cash flow", "high_dti": "High debt-to-income ratio: {}%",
      "low_save": "Low savings rate: {}%", "low_emerg": "Emergency fund below 3 months ({} months)",
      "low_goal": "Low feasibility for '{}': {}%", "bad_goal": "Incomplete goal ignored: {}",
      "contradiction": "Current savings exceed target: {}", "debt_mismatch": "Debt payment entered but no debt balance"}


# ---------------- Calculation engine (deterministic, no LLM) ----------------
def num(v, d=0.0):
    try:
        return float(v)
    except (TypeError, ValueError):
        return d


def pmt(gap, r, n):
    if gap <= 0 or n <= 0:
        return 0.0
    return gap * r / ((1 + r) ** n - 1) if r else gap / n


def core(p):
    inc = p["income"] * (1 + p["income_adj"] / 100)
    exp = p["expenses"] * (1 + p["expense_adj"] / 100)
    surplus = inc - exp - p["debt_payment"]
    budget = max(0, surplus) * p["save_pct"] / 100
    r = p["return_rate"] / 1200
    gs = []
    for g in p["goals"]:
        yrs = max(g["years"] + p["horizon_adj"], 0.25)
        n = yrs * 12
        req = pmt(max(0, g["target"] - g["current"] * (1 + r) ** n), r, n)
        solo = min(100.0, budget / req * 100) if req else 100.0
        score = 0.5 * max(0, 100 - yrs * 8) + 0.3 * TYPE_W.get(g["type"], 40) + 0.2 * solo
        gs.append(dict(name=g["name"], type=g["type"], target=g["target"], years=round(yrs, 2),
                       required=req, score=score, standalone=solo))
    left = budget
    for rank, i in enumerate(sorted(range(len(gs)), key=lambda i: -gs[i]["score"]), 1):
        g = gs[i]
        a = min(g["required"], left)
        left -= a
        g.update(rank=rank, funded=a, feasibility=a / g["required"] * 100 if g["required"] else 100.0,
                 reason="type" if TYPE_W.get(g["type"], 40) >= 70 else ("urgent" if g["years"] <= 3 else "afford"))
    tot = sum(g["required"] for g in gs)
    alloc = sum(g["funded"] for g in gs)
    return dict(income=inc, expenses=exp, debt_payment=p["debt_payment"], surplus=surplus, budget=budget,
                goals=gs, total_required=tot, allocated=alloc, overall=alloc / tot * 100 if tot else 100.0)


def normalize(d):
    goals = [dict(name=str(x.get("name") or "Goal")[:60], type=x.get("type") or "other", target=num(x.get("target")),
                  current=num(x.get("current")), years=num(x.get("years"))) for x in (d.get("goals") or [])[:3]]
    p = {k: num(d.get(k)) for k in ("income", "expenses", "debt", "debt_payment", "cash", "investments",
                                    "other_assets", "income_adj", "expense_adj", "horizon_adj")}
    p.update(return_rate=num(d.get("return_rate"), 6), save_pct=num(d.get("save_pct"), 100), goals=goals)
    return p


def w(code, *v):
    return {"code": code, "v": list(v)}


def analyze(d):
    p = normalize(d)
    warns = []
    if p["income"] <= 0:
        return {"error": "no_income"}
    for g in p["goals"]:
        if g["target"] <= 0 or g["years"] <= 0:
            warns.append(w("bad_goal", g["name"]))
        elif g["current"] > g["target"]:
            warns.append(w("contradiction", g["name"]))
    if p["debt_payment"] > 0 and p["debt"] <= 0:
        warns.append(w("debt_mismatch"))
    p["goals"] = [g for g in p["goals"] if g["target"] > 0 and g["years"] > 0]
    if not p["goals"]:
        return {"error": "no_goals"}
    base = core(p)
    m = dict(save_rate=base["surplus"] / base["income"] * 100, dti=p["debt_payment"] / base["income"] * 100,
             emerg=p["cash"] / base["expenses"] if base["expenses"] else 99,
             net_worth=p["cash"] + p["investments"] + p["other_assets"] - p["debt"])
    if base["surplus"] < 0:
        warns.append(w("neg_cf"))
    if m["dti"] > 40:
        warns.append(w("high_dti", round(m["dti"], 1)))
    if 0 <= m["save_rate"] < 10:
        warns.append(w("low_save", round(m["save_rate"], 1)))
    if m["emerg"] < 3:
        warns.append(w("low_emerg", round(m["emerg"], 1)))
    for g in base["goals"]:
        if g["feasibility"] < 50:
            warns.append(w("low_goal", g["name"], round(g["feasibility"])))
    r0 = p["return_rate"]
    scen = []
    for code, dr in (("sc_cons", -2), ("sc_base", 0), ("sc_opt", 2)):
        c = core({**p, "return_rate": max(0, r0 + dr)})
        scen.append(dict(code=code, rate=max(0, r0 + dr), total_required=c["total_required"], overall=c["overall"]))
    ia, ea = p["income_adj"], p["expense_adj"]
    cases = [("s_base", {}), ("s_inc", {"income_adj": ia - 20}), ("s_exp", {"expense_adj": ea + 15}),
             ("s_both", {"income_adj": ia - 20, "expense_adj": ea + 15}), ("s_ret", {"return_rate": max(0, r0 - 3)})]
    stress = []
    for code, ch in cases:
        c = core({**p, **ch})
        stress.append(dict(code=code, surplus=c["surplus"], overall=c["overall"],
                           goals=[g["feasibility"] for g in c["goals"]]))
    return dict(base=base, metrics=m, scenarios=scen, stress=stress, warnings=warns,
                assumptions={k: p[k] for k in ("return_rate", "save_pct", "income_adj", "expense_adj", "horizon_adj")})


# ---------------- Gemini (Google AI Studio) ----------------
def ask(prompt, key, js=False):
    key = key or os.environ.get("GEMINI_API_KEY", "")
    if not key:
        raise ValueError("no_key")
    cfg = {"temperature": 0.3}
    if js:
        cfg["responseMimeType"] = "application/json"
    err = ""
    for m in [x for x in (os.environ.get("GEMINI_MODEL"), "gemini-3.8-flash", "gemini-3.6-flash") if x]:
        r = requests.post(f"https://generativelanguage.googleapis.com/v1beta/models/{m}:generateContent",
                          headers={"x-goog-api-key": key},
                          json={"contents": [{"parts": [{"text": prompt}]}], "generationConfig": cfg}, timeout=60)
        if r.ok:
            try:
                return "".join(x.get("text", "") for x in r.json()["candidates"][0]["content"]["parts"])
            except Exception:
                err = "empty response"
        else:
            err = r.text[:200]
    raise RuntimeError(err)


def lang_name(l):
    return "Vietnamese" if l == "vi" else "English"


@app.post("/api/intake")
def intake():
    j = request.json or {}
    prompt = f"""Role: data-extraction assistant for a personal finance planner.
Task: extract the customer's financial profile from the text. Amounts in VND per month unless a total is natural (debt, cash, investments, other_assets, target, current). Convert 'triệu'=1,000,000 and 'tỷ'=1,000,000,000.
Format: JSON only: {{"income":n|null,"expenses":n|null,"debt":n|null,"debt_payment":n|null,"cash":n|null,"investments":n|null,"other_assets":n|null,"goals":[{{"name":s,"type":"emergency|home|education|retirement|other","target":n,"current":n|null,"years":n}}],"missing":[names of important fields not stated],"questions":[short clarifying questions in {lang_name(j.get('lang'))}]}}
Rules: use null when not stated; never invent numbers; max 3 goals.
TEXT: {str(j.get('text', ''))[:4000]}"""
    try:
        t = ask(prompt, j.get("key"), js=True)
        m = re.search(r"\{.*\}", t, re.S)
        return jsonify(json.loads(m.group(0)))
    except Exception as e:
        return jsonify({"error": str(e)}), 400


@app.post("/api/advice")
def advice():
    j = request.json or {}
    if not j.get("confirmed"):
        return jsonify({"error": "assumptions not confirmed"}), 400
    res = j.get("result") or {}
    prompt = f"""Role: careful personal-finance planning educator in Vietnam.
Context: all figures and calculations below (VND) were computed by a Python engine. Use ONLY these numbers; do not invent or recompute. If data is missing or contradictory, say so and ask the user to confirm.
Task: (1) explain why the goal with rank 1 is prioritised and the trade-offs against the other goals; (2) give personalised planning advice on cash flow, debt, emergency fund and goal timing; (3) explain the stress-test results and main risks.
Format: Markdown with sections 'Priority rationale', 'Personalised advice', 'Stress test & risks', 'Questions to confirm'; at most 350 words.
Guardrails: educational only; name no specific stocks, funds or products; no guaranteed returns; state uncertainty. Write in {lang_name(j.get('lang'))}.
DATA: {json.dumps({'profile': j.get('data'), 'result': res}, ensure_ascii=False)[:12000]}"""
    try:
        return jsonify({"text": ask(prompt, j.get("key"))})
    except ValueError:
        return jsonify({"error": "no_key"}), 400
    except Exception as e:
        return jsonify({"error": str(e)}), 502


@app.post("/api/analyze")
def api_analyze():
    return jsonify(analyze(request.json or {}))


# ---------------- Storage & Excel ----------------
def xlsx(sheets):
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name[:31])
        for r in rows:
            ws.append(r)
        for c in ws[1]:
            c.font = Font(bold=True)
        ws.column_dimensions["A"].width = 34
        ws.column_dimensions["B"].width = 20
    b = io.BytesIO()
    wb.save(b)
    b.seek(0)
    return send_file(b, as_attachment=True, download_name="export.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.post("/api/save")
def save():
    j = request.json or {}
    if not j.get("confirmed") or not j.get("result") or "error" in j["result"]:
        return jsonify({"error": "nothing to save"}), 400
    with engine.begin() as c:
        rid = c.execute(insert(plans).values(created=datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S"),
                        name=str(j.get("name") or "Anonymous")[:200],
                        payload=json.dumps({"data": j.get("data"), "result": j["result"], "advice": j.get("advice", "")}))
                        ).inserted_primary_key[0]
    return jsonify({"id": rid})


@app.post("/api/export")
def export():
    j = request.json or {}
    d, r = j.get("data") or {}, j.get("result") or {}
    if not r or "error" in r:
        return jsonify({"error": "no result"}), 400
    b, m, a = r["base"], r["metrics"], r["assumptions"]
    return xlsx({
        "Profile": [["Item", "Value"]] + [[k, d.get(k, 0)] for k in ("income", "expenses", "debt", "debt_payment", "cash", "investments", "other_assets")]
                   + [["(VND)", ""]] + [[k, v] for k, v in a.items()],
        "Analysis": [["Metric", "Value"], ["Monthly income (VND)", b["income"]], ["Monthly expenses (VND)", b["expenses"]],
                     ["Monthly surplus (VND)", b["surplus"]], ["Savings rate (%)", m["save_rate"]], ["Debt-to-income (%)", m["dti"]],
                     ["Emergency fund (months)", m["emerg"]], ["Net worth (VND)", m["net_worth"]], ["Goals funded overall (%)", b["overall"]]],
        "Goals": [["Rank", "Goal", "Target (VND)", "Years", "Required / month (VND)", "Funded / month (VND)", "Feasibility (%)"]]
                 + [[g["rank"], g["name"], g["target"], g["years"], g["required"], g["funded"], g["feasibility"]] for g in sorted(b["goals"], key=lambda g: g["rank"])],
        "Scenarios": [["Scenario", "Return (%)", "Total required / month (VND)", "Funded overall (%)"]] + [[s["code"], s["rate"], s["total_required"], s["overall"]] for s in r["scenarios"]],
        "Stress Test": [["Case", "Surplus (VND)", "Funded overall (%)"]] + [[s["code"], s["surplus"], s["overall"]] for s in r["stress"]],
        "Warnings": [["Warning"]] + [[WE[x["code"]].format(*x["v"])] for x in r["warnings"]],
        "AI Advice": [["Personalised advice"]] + [[l] for l in (j.get("advice") or "").split("\n") if l.strip()],
    })


@app.post("/api/admin/export")
def admin_export():
    k = os.environ.get("ADMIN_KEY")
    if not k or request.headers.get("X-Admin-Key") != k:
        return jsonify({"error": "unauthorized"}), 401
    cust, goals = [["ID", "Created (UTC)", "Name", "Income", "Expenses", "Surplus", "Funded overall (%)", "Advice saved"]], \
                  [["Customer ID", "Goal", "Target (VND)", "Years", "Required / month", "Feasibility (%)", "Rank"]]
    with engine.connect() as c:
        for row in c.execute(select(plans).order_by(plans.c.id)):
            p = json.loads(row.payload)
            b = p["result"]["base"]
            cust.append([row.id, row.created, row.name, b["income"], b["expenses"], b["surplus"], b["overall"], "yes" if p.get("advice") else "no"])
            goals += [[row.id, g["name"], g["target"], g["years"], g["required"], g["feasibility"], g["rank"]] for g in b["goals"]]
    return xlsx({"Customers": cust, "Goals": goals})


@app.get("/")
def index():
    return app.send_static_file("index.html")


if __name__ == "__main__":
    app.run(debug=True, port=int(os.environ.get("PORT", 5000)))
